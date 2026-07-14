from __future__ import annotations

import calendar
from datetime import date, datetime
from io import BytesIO
import re
from typing import Any

from openpyxl import load_workbook

from .base import SourceAdapter
from ..models import CollectionResult, Record


class CERElectricityTradeAdapter(SourceAdapter):
    """Collect recurring CER electricity-trade workbook observations.

    The workbook contains national monthly volumes/values and western/eastern
    regional price series. It does not identify B.C. imports separately, so
    geography is retained exactly as published rather than inferred.
    """

    def collect(self, full_refresh: bool = False) -> CollectionResult:
        response = self.http.request_bytes(
            self.source_id,
            "GET",
            self.source["download_url"],
            label="electricity-trade-summary",
        )
        result = CollectionResult(self.source_id, artifacts=[response.artifact])
        workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
        publication_date = self._publication_date(workbook)
        self._collect_monthly_trade(workbook, response.artifact.document_id, publication_date, result)
        self._collect_monthly_prices(workbook, response.artifact.document_id, publication_date, result)
        workbook.close()
        return result

    def _collect_monthly_trade(
        self,
        workbook: Any,
        document_id: str,
        publication_date: str | None,
        result: CollectionResult,
    ) -> None:
        sheet_name = self.source.get("monthly_trade_sheet", "Fig. 1(m), Fig. 3(m)")
        sheet = workbook[sheet_name]
        headers = [self._text(value) for value in next(sheet.iter_rows(values_only=True))]
        series = self.source.get("monthly_trade_series", [])
        columns = self._match_columns(headers, series, result)
        start_limit = date.fromisoformat(str(self.source.get("start_date", "2010-01-01"))[:10])
        for values in sheet.iter_rows(min_row=2, values_only=True):
            period = self._month(values[0] if values else None)
            if period is None or period < start_limit:
                continue
            end = date(period.year, period.month, calendar.monthrange(period.year, period.month)[1])
            for index, config in columns.items():
                value = self._number(values[index] if index < len(values) else None)
                if value is None:
                    continue
                result.records.append(
                    Record(
                        variable_id=config["variable_id"],
                        source_id=self.source_id,
                        source_series_id=config["source_series_id"],
                        value=value,
                        unit_original=config["unit_original"],
                        unit_canonical=config["unit_canonical"],
                        reference_period_start=period.isoformat(),
                        reference_period_end=end.isoformat(),
                        period_basis="monthly",
                        publication_date=publication_date,
                        geography_id=config.get("geography_id"),
                        currency=config.get("currency"),
                        document_id=document_id,
                        extraction_method="xlsx",
                        metadata={"workbook_sheet": sheet_name, "column_header": headers[index]},
                    )
                )

    def _collect_monthly_prices(
        self,
        workbook: Any,
        document_id: str,
        publication_date: str | None,
        result: CollectionResult,
    ) -> None:
        sheet_name = self.source.get("monthly_price_sheet", "Fig. 4")
        sheet = workbook[sheet_name]
        headers = [self._text(value) for value in next(sheet.iter_rows(values_only=True))]
        series = self.source.get("monthly_price_series", [])
        columns = self._match_columns(headers, series, result)
        start_limit = date.fromisoformat(str(self.source.get("start_date", "2010-01-01"))[:10])
        for values in sheet.iter_rows(min_row=2, values_only=True):
            period = self._month(values[0] if values else None)
            if period is None or period < start_limit:
                continue
            end = date(period.year, period.month, calendar.monthrange(period.year, period.month)[1])
            for index, config in columns.items():
                value = self._number(values[index] if index < len(values) else None)
                if value is None:
                    continue
                result.records.append(
                    Record(
                        variable_id=config["variable_id"],
                        source_id=self.source_id,
                        source_series_id=config["source_series_id"],
                        value=value,
                        unit_original=config["unit_original"],
                        unit_canonical=config["unit_canonical"],
                        reference_period_start=period.isoformat(),
                        reference_period_end=end.isoformat(),
                        period_basis="monthly",
                        publication_date=publication_date,
                        geography_id=config.get("geography_id"),
                        currency=config.get("currency", "CAD"),
                        document_id=document_id,
                        extraction_method="xlsx",
                        metadata={"workbook_sheet": sheet_name, "column_header": headers[index]},
                    )
                )

    @staticmethod
    def _match_columns(headers: list[str], series: list[dict[str, Any]], result: CollectionResult) -> dict[int, dict[str, Any]]:
        matched: dict[int, dict[str, Any]] = {}
        for config in series:
            patterns = config.get("header_patterns", [])
            indices = [
                index
                for index, header in enumerate(headers)
                if header and any(re.search(pattern, header, re.IGNORECASE) for pattern in patterns)
            ]
            if len(indices) != 1:
                result.warnings.append(
                    f"CER series {config['source_series_id']!r} matched {len(indices)} columns"
                )
                continue
            matched[indices[0]] = config
        return matched

    @staticmethod
    def _publication_date(workbook: Any) -> str | None:
        if "Source Info" not in workbook.sheetnames:
            return None
        sheet = workbook["Source Info"]
        for row in sheet.iter_rows(min_row=1, max_row=12, values_only=True):
            for value in row:
                text = str(value or "")
                match = re.search(r"Updated\s+(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})", text, re.I)
                if match:
                    try:
                        return datetime.strptime(" ".join(match.groups()), "%d %B %Y").date().isoformat()
                    except ValueError:
                        return None
        return None

    @staticmethod
    def _month(value: Any) -> date | None:
        if isinstance(value, datetime):
            return date(value.year, value.month, 1)
        if isinstance(value, date):
            return date(value.year, value.month, 1)
        return None

    @staticmethod
    def _number(value: Any) -> float | None:
        if value is None or str(value).strip() in {"", "-", "..", "N/A"}:
            return None
        try:
            return float(str(value).replace(",", ""))
        except ValueError:
            return None

    @staticmethod
    def _text(value: Any) -> str:
        return re.sub(r"\s+", " ", str(value or "").replace("\n", " ")).strip()
