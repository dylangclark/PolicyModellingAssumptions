from io import BytesIO

from openpyxl import Workbook

from bc_assumptions.adapters.cer_electricity_trade import CERElectricityTradeAdapter
from bc_assumptions.http import BytesResult
from bc_assumptions.models import Artifact


class FakeHttp:
    def __init__(self, content, artifact):
        self.content = content
        self.artifact = artifact

    def request_bytes(self, *args, **kwargs):
        return BytesResult(self.artifact, self.content)


def workbook_bytes():
    workbook = Workbook()
    info = workbook.active
    info.title = "Source Info"
    info.append(["Updated 16 June 2026"])
    trade = workbook.create_sheet("Fig. 1(m), Fig. 3(m)")
    trade.append(["Date", "Exports / Exportations (MW.h)", "Imports / Importations (MW.h)", "Exports / Exportations ($)", "Imports / Importations ($)"])
    trade.append([__import__('datetime').datetime(2026, 5, 1), 10, 20, 30, 40])
    prices = workbook.create_sheet("Fig. 4")
    prices.append(["Date", "Exports from Canada – West ($/MW.h)", "Imports to Canada – West ($/MW.h)"])
    prices.append([__import__('datetime').datetime(2026, 5, 1), 50, 60])
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_cer_adapter(db):
    artifact = Artifact("cer", "https://example.test/a.xlsx", "cer/a.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "2026-07-14T00:00:00Z", "a"*64, 100)
    source = {
        "id": "cer", "download_url": "https://example.test/a.xlsx", "start_date": "2010-01-01",
        "monthly_trade_series": [
            {"source_series_id":"exports","variable_id":"exports","header_patterns":["^Exports.*MW\\.h"],"unit_original":"MWh","unit_canonical":"MWh","geography_id":"CA"},
            {"source_series_id":"imports","variable_id":"imports","header_patterns":["^Imports.*MW\\.h"],"unit_original":"MWh","unit_canonical":"MWh","geography_id":"CA"},
        ],
        "monthly_price_series": [
            {"source_series_id":"west_export","variable_id":"west_export","header_patterns":["Exports from Canada.*West.*\\$/MW\\.h"],"unit_original":"CAD_per_MWh","unit_canonical":"CAD_per_MWh","geography_id":"CA-WEST"},
            {"source_series_id":"west_import","variable_id":"west_import","header_patterns":["Imports to Canada.*West.*\\$/MW\\.h"],"unit_original":"CAD_per_MWh","unit_canonical":"CAD_per_MWh","geography_id":"CA-WEST"},
        ],
    }
    variables = {key: {"geography_id": "CA"} for key in ["exports","imports","west_export","west_import"]}
    result = CERElectricityTradeAdapter(source, variables, db, FakeHttp(workbook_bytes(), artifact)).collect()
    assert len(result.records) == 4
    assert {r.value for r in result.records} == {10.0, 20.0, 50.0, 60.0}
    assert all(r.publication_date == "2026-06-16" for r in result.records)
