from __future__ import annotations

import calendar
from datetime import date, datetime
from io import BytesIO
import re
from typing import Any

from openpyxl import load_workbook

from .base import SourceAdapter
from ..models import CollectionResult, Record


class WorldBankPinkSheetAdapter(SourceAdapter):
    def collect(self, full_refresh: bool = False) -> CollectionResult:
        response = self.http.request_bytes(
            self.source_id,
            "GET",
            self.source["download_url"],
            label="monthly-prices",
        )
        result = CollectionResult(self.source_id, artifacts=[response.artifact])
        workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
        sheet = self._monthly_sheet(workbook, self.source.get("sheet_regex"))
        header_row, headers = self._find_headers(sheet)
        matched_columns: dict[int, dict[str, Any]] = {}
        for config in self.source.get("series", []):
            matches = [
                index
                for index, header in enumerate(headers)
                if header
                and any(
                    re.search(pattern, header, re.IGNORECASE)
                    for pattern in config.get("aliases", [])
                )
            ]
            if len(matches) == 1:
                matched_columns[matches[0]] = config
            else:
                result.warnings.append(
                    f"World Bank series {config['source_series_id']!r} matched {len(matches)} columns"
                )

        start_limit = date.fromisoformat(str(self.source.get("start_date", "2000-01-01"))[:10])
        for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            if not values:
                continue
            period = self._parse_month(values[0])
            if not period or period < start_limit:
                continue
            end = date(period.year, period.month, calendar.monthrange(period.year, period.month)[1])
            for index, config in matched_columns.items():
                if index >= len(values):
                    continue
                value = self._number(values[index])
                if value is None:
                    continue
                variable = self.variables[config["variable_id"]]
                flags = list(config.get("quality_flags", []))
                if "proxy" in config["variable_id"] and "proxy_series" not in flags:
                    flags.append("proxy_series")
                result.records.append(
                    Record(
                        variable_id=config["variable_id"],
                        source_id=self.source_id,
                        source_series_id=config["source_series_id"],
                        value=value,
                        unit_original=config["unit_original"],
                        unit_canonical=config["unit_canonical"],
                        reference_period_start=period.isoformat(),
                        reference_period_end=end.isoformat(),
                        period_basis="monthly",
                        geography_id=variable.get("geography_id"),
                        document_id=response.artifact.document_id,
                        extraction_method="xlsx",
                        quality_flags=flags,
                        metadata={
                            "workbook_sheet": sheet.title,
                            "column_header": headers[index],
                            "is_jkm": False
                            if "lng" in config["source_series_id"].lower()
                            and "japan" in config["source_series_id"].lower()
                            else None,
                        },
                    )
                )
        workbook.close()
        return result

    @staticmethod
    def _monthly_sheet(workbook: Any, sheet_regex: str | None = None) -> Any:
        if sheet_regex:
            matches = [
                name
                for name in workbook.sheetnames
                if re.search(sheet_regex, name, flags=re.IGNORECASE)
            ]
            if len(matches) != 1:
                raise RuntimeError(
                    f"World Bank sheet selector {sheet_regex!r} matched {len(matches)} sheets: "
                    f"{workbook.sheetnames}"
                )
            return workbook[matches[0]]
        for name in workbook.sheetnames:
            normalized = re.sub(r"[^a-z]", "", name.lower())
            if "monthly" in normalized:
                return workbook[name]
        return workbook[workbook.sheetnames[0]]

    @staticmethod
    def _find_headers(sheet: Any) -> tuple[int, list[str]]:
        rows = list(sheet.iter_rows(min_row=1, max_row=30, values_only=True))
        for row_number, values in enumerate(rows, start=1):
            normalized = [WorldBankPinkSheetAdapter._header(value) for value in values]
            if (
                normalized
                and normalized[0].lower() == "date"
                and sum(bool(value) for value in normalized) >= 2
            ):
                return row_number, normalized

            # Current Pink Sheet workbooks leave the period header blank. Confirm a
            # candidate commodity-name row by finding a monthly period shortly below it;
            # this avoids mistaking the title or units rows for the header.
            following = rows[row_number : row_number + 5]
            if (
                normalized
                and not normalized[0]
                and sum(bool(value) for value in normalized) >= 2
                and any(
                    later and WorldBankPinkSheetAdapter._parse_month(later[0]) is not None
                    for later in following
                )
            ):
                return row_number, normalized
        raise RuntimeError("Could not find the commodity header in the World Bank workbook")

    @staticmethod
    def _header(value: Any) -> str:
        if value is None:
            return ""
        return re.sub(r"\s+", " ", str(value).replace("\n", " ")).strip()

    @staticmethod
    def _parse_month(value: Any) -> date | None:
        if isinstance(value, datetime):
            return date(value.year, value.month, 1)
        if isinstance(value, date):
            return date(value.year, value.month, 1)
        text = str(value or "").strip()
        for pattern in [r"^(\d{4})M(\d{1,2})$", r"^(\d{4})[-/](\d{1,2})$"]:
            match = re.match(pattern, text, flags=re.IGNORECASE)
            if match:
                return date(int(match.group(1)), int(match.group(2)), 1)
        try:
            parsed = datetime.fromisoformat(text)
            return date(parsed.year, parsed.month, 1)
        except ValueError:
            return None

    @staticmethod
    def _number(value: Any) -> float | None:
        if value is None or str(value).strip() in {"", "..", "-", "NA", "N/A"}:
            return None
        try:
            return float(str(value).replace(",", ""))
        except ValueError:
            return None
